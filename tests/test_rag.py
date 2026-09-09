from datetime import datetime, timezone
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy.dialects import postgresql

from app.ai.loaders import (
    DocumentTooLargeError,
    extract_document_sections,
)
from app.ai.text_splitter import split_text
from app.core.config import settings
from app.core.exceptions import ClassFlowError
from app.models.resource import (
    RAG_SOURCE_RESOURCE,
    RAG_SOURCE_TASK,
    RAG_SOURCE_TASK_ATTACHMENT,
    RagChunk,
)
from app.models.task import (
    TASK_STATUS_ACTIVE,
    TASK_VISIBILITY_PERSONAL,
    TASK_VISIBILITY_SHARED,
    Task,
    TaskAttachment,
)
from app.repositories.rag import RagRepository, RagSearchResult
from app.services.rag import NO_CONTEXT_ANSWER, RagChatService


class FakeRagRepository:
    def __init__(
        self,
        matches: list[RagSearchResult],
        has_class_access: bool = True,
    ) -> None:
        self.matches = matches
        self.has_class_access = has_class_access
        self.search_arguments: dict | None = None
        self.replace_arguments: dict | None = None
        self.deleted_source: tuple[str, int] | None = None

    async def user_has_approved_class_access(
        self,
        classroom_id: int,
        user_id: int,
    ) -> bool:
        return self.has_class_access

    async def search_accessible(self, **kwargs) -> list[RagSearchResult]:
        self.search_arguments = kwargs
        return self.matches

    async def replace_source_chunks(self, **kwargs) -> list[RagChunk]:
        self.replace_arguments = kwargs
        return []

    async def delete_source_chunks(self, source_type: str, source_id: int) -> None:
        self.deleted_source = (source_type, source_id)


class FakeGeminiClient:
    def __init__(self, answer: str = "The project requires a report [1].") -> None:
        self.answer = answer
        self.generated_contexts: list[str] | None = None
        self.embedded_documents: list[str] | None = None

    async def embed_query(self, text: str) -> list[float]:
        return [0.1] * 768

    async def generate_answer(self, question: str, contexts: list[str]) -> str:
        self.generated_contexts = contexts
        return self.answer

    async def embed_documents(self, texts: list[str], title: str | None = None) -> list[list[float]]:
        self.embedded_documents = texts
        return [[0.1] * 768 for _ in texts]


class FakeSession:
    def __init__(self) -> None:
        self.commit_count = 0

    async def commit(self) -> None:
        self.commit_count += 1


def make_chunk(
    chunk_id: int,
    source_type: str,
    source_id: int,
    content: str,
) -> RagChunk:
    return RagChunk(
        id=chunk_id,
        classroom_id=10,
        class_course_id=None,
        source_type=source_type,
        source_id=source_id,
        chunk_index=0,
        content=content,
        embedding=[0.1] * 768,
    )


def test_split_text_normalizes_and_overlaps_chunks() -> None:
    chunks = split_text(
        "Alpha   beta gamma delta epsilon zeta eta theta",
        chunk_size=24,
        chunk_overlap=6,
    )

    assert len(chunks) > 1
    assert all(len(chunk) <= 24 for chunk in chunks)
    assert all("  " not in chunk for chunk in chunks)
    assert chunks[0] == "Alpha beta gamma delta"


@pytest.mark.anyio
async def test_chat_returns_safe_response_when_no_chunks_are_accessible() -> None:
    repository = FakeRagRepository([])
    ai_client = FakeGeminiClient()
    service = RagChatService(
        session=object(),
        repository=repository,
        ai_client=ai_client,
    )

    response = await service.answer_class_question(10, 22, "What is due?")

    assert response.answer == NO_CONTEXT_ANSWER
    assert response.sources == []
    assert ai_client.generated_contexts is None
    assert repository.search_arguments is not None
    assert repository.search_arguments["classroom_id"] == 10
    assert repository.search_arguments["user_id"] == 22


@pytest.mark.anyio
async def test_chat_returns_only_sources_cited_by_the_answer() -> None:
    task_chunk = make_chunk(1, RAG_SOURCE_TASK, 30, "Submit a report before Friday.")
    second_task_chunk = make_chunk(2, RAG_SOURCE_TASK, 30, "The report must be a PDF.")
    resource_chunk = make_chunk(3, RAG_SOURCE_RESOURCE, 40, "Use the supplied template.")
    repository = FakeRagRepository(
        [
            RagSearchResult(task_chunk, 0.1),
            RagSearchResult(second_task_chunk, 0.2),
            RagSearchResult(resource_chunk, 0.3),
        ]
    )
    ai_client = FakeGeminiClient()
    service = RagChatService(
        session=object(),
        repository=repository,
        ai_client=ai_client,
    )

    response = await service.answer_class_question(10, 22, "What is required?")

    assert response.answer == "The project requires a report [1]."
    assert [(source.source_type, source.source_id) for source in response.sources] == [
        (RAG_SOURCE_TASK, 30),
    ]
    assert ai_client.generated_contexts == [
        task_chunk.content,
        second_task_chunk.content,
        resource_chunk.content,
    ]


@pytest.mark.anyio
async def test_chat_supports_grouped_citations_and_deduplicates_sources() -> None:
    first_task_chunk = make_chunk(1, RAG_SOURCE_TASK, 30, "Submit a report.")
    second_task_chunk = make_chunk(2, RAG_SOURCE_TASK, 30, "The report must be a PDF.")
    resource_chunk = make_chunk(3, RAG_SOURCE_RESOURCE, 40, "Use the template.")
    repository = FakeRagRepository(
        [
            RagSearchResult(first_task_chunk, 0.1),
            RagSearchResult(second_task_chunk, 0.2),
            RagSearchResult(resource_chunk, 0.3),
        ]
    )
    ai_client = FakeGeminiClient(
        "Submit one PDF report using the supplied template [1, 2, 3]."
    )
    service = RagChatService(
        session=object(),
        repository=repository,
        ai_client=ai_client,
    )

    response = await service.answer_class_question(10, 22, "What is required?")

    assert [(source.source_type, source.source_id) for source in response.sources] == [
        (RAG_SOURCE_TASK, 30),
        (RAG_SOURCE_RESOURCE, 40),
    ]


@pytest.mark.anyio
async def test_chat_ignores_invalid_or_missing_source_citations() -> None:
    task_chunk = make_chunk(1, RAG_SOURCE_TASK, 30, "Submit a report.")
    repository = FakeRagRepository([RagSearchResult(task_chunk, 0.1)])
    ai_client = FakeGeminiClient("The information is unavailable [9].")
    service = RagChatService(
        session=object(),
        repository=repository,
        ai_client=ai_client,
    )

    response = await service.answer_class_question(10, 22, "What is required?")

    assert response.sources == []


def test_retrieval_query_enforces_membership_course_and_source_policy() -> None:
    repository = RagRepository(session=object())
    statement = repository._accessible_search_statement(
        classroom_id=10,
        user_id=22,
        query_embedding=[0.1] * 768,
        limit=5,
        max_distance=0.8,
    )

    compiled = str(statement.compile(dialect=postgresql.dialect()))

    assert "class_memberships.status" in compiled
    assert "class_memberships.user_id" in compiled
    assert "course_registrations.is_active IS true" in compiled
    assert "class_courses.is_active IS true" in compiled
    assert "tasks.visibility" in compiled
    assert "tasks.status" in compiled
    assert "resources.is_enabled IS true" in compiled
    assert "task_attachments" in compiled
    assert "rag_chunks.classroom_id" in compiled


@pytest.mark.anyio
async def test_index_task_builds_embeddings_for_active_shared_task() -> None:
    repository = FakeRagRepository([])
    ai_client = FakeGeminiClient()
    session = FakeSession()
    service = RagChatService(session=session, repository=repository, ai_client=ai_client)
    task = Task(
        id=11,
        classroom_id=1,
        class_course_id=2,
        created_by_user_id=3,
        title="Assignment 3",
        description="Submit the implementation.",
        task_type="assignment",
        visibility=TASK_VISIBILITY_SHARED,
        priority="high",
        status=TASK_STATUS_ACTIVE,
        deadline=datetime(2026, 8, 28, 14, 7, tzinfo=timezone.utc),
    )

    indexed_count = await service.index_task(task)

    assert indexed_count == 1
    assert ai_client.embedded_documents is not None
    assert "Task title: Assignment 3" in ai_client.embedded_documents[0]
    assert "Deadline: 2026-08-28T14:07:00+00:00" in ai_client.embedded_documents[0]
    assert repository.replace_arguments is not None
    assert repository.replace_arguments["source_type"] == RAG_SOURCE_TASK
    assert repository.replace_arguments["source_id"] == 11
    assert session.commit_count == 1


@pytest.mark.anyio
async def test_index_task_never_embeds_personal_task() -> None:
    repository = FakeRagRepository([])
    ai_client = FakeGeminiClient()
    session = FakeSession()
    service = RagChatService(session=session, repository=repository, ai_client=ai_client)
    task = Task(
        id=12,
        classroom_id=1,
        class_course_id=None,
        created_by_user_id=3,
        title="Private study plan",
        description="Personal notes",
        task_type="other",
        visibility=TASK_VISIBILITY_PERSONAL,
        priority="medium",
        status=TASK_STATUS_ACTIVE,
        deadline=None,
    )

    indexed_count = await service.index_task(task)

    assert indexed_count == 0
    assert ai_client.embedded_documents is None
    assert repository.deleted_source == (RAG_SOURCE_TASK, 12)
    assert session.commit_count == 1


def test_extract_docx_includes_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("The assignment requires a UML diagram.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Deliverable"
    table.cell(0, 1).text = "Source code"
    buffer = BytesIO()
    document.save(buffer)

    sections = extract_document_sections("requirements.docx", buffer.getvalue())

    assert len(sections) == 1
    assert "The assignment requires a UML diagram." in sections[0].text
    assert "Deliverable | Source code" in sections[0].text
    assert sections[0].page_number is None


def test_extract_pdf_preserves_page_numbers(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        is_encrypted = False
        pages = [FakePage("First page"), FakePage("Second page")]

    monkeypatch.setattr("app.ai.loaders.PdfReader", lambda *_args, **_kwargs: FakeReader())

    sections = extract_document_sections("outline.pdf", b"%PDF-test", max_pages=2)

    assert [(section.page_number, section.text) for section in sections] == [
        (1, "First page"),
        (2, "Second page"),
    ]


def test_extract_pdf_rejects_document_above_page_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeReader:
        is_encrypted = False
        pages = [object(), object(), object()]

    monkeypatch.setattr("app.ai.loaders.PdfReader", lambda *_args, **_kwargs: FakeReader())

    with pytest.raises(DocumentTooLargeError):
        extract_document_sections("large.pdf", b"%PDF-test", max_pages=2)


@pytest.mark.anyio
async def test_index_task_attachment_preserves_file_source_metadata() -> None:
    repository = FakeRagRepository([])
    ai_client = FakeGeminiClient()
    session = FakeSession()
    service = RagChatService(session=session, repository=repository, ai_client=ai_client)
    task = Task(
        id=18,
        classroom_id=1,
        class_course_id=3,
        created_by_user_id=3,
        title="Assignment 4",
        description=None,
        task_type="assignment",
        visibility=TASK_VISIBILITY_SHARED,
        priority="medium",
        status=TASK_STATUS_ACTIVE,
        deadline=None,
    )
    attachment = TaskAttachment(
        id=5,
        task_id=18,
        uploaded_by_user_id=3,
        file_name="requirements.txt",
        storage_key="tasks/18/requirements.txt",
        file_type="text/plain",
        file_size=41,
    )

    indexed_count = await service.index_task_attachment(
        task,
        attachment,
        b"Implement inheritance and submit a report.",
    )

    assert indexed_count == 1
    assert ai_client.embedded_documents is not None
    assert "Attachment: requirements.txt" in ai_client.embedded_documents[0]
    assert "Implement inheritance" in ai_client.embedded_documents[0]
    assert repository.replace_arguments is not None
    assert repository.replace_arguments["source_type"] == RAG_SOURCE_TASK_ATTACHMENT
    assert repository.replace_arguments["source_id"] == 5
    assert repository.replace_arguments["source_title"] == "requirements.txt"
    assert repository.replace_arguments["page_numbers"] == [None]
    assert session.commit_count == 1


@pytest.mark.anyio
async def test_chat_with_docx_uses_temporary_chunks_without_persisting_them() -> None:
    document = Document()
    document.add_paragraph(
        "The capstone report must include a UML class diagram and testing evidence."
    )
    buffer = BytesIO()
    document.save(buffer)
    repository = FakeRagRepository([])
    ai_client = FakeGeminiClient()
    service = RagChatService(
        session=object(),
        repository=repository,
        ai_client=ai_client,
    )

    response = await service.answer_class_question_with_file(
        classroom_id=10,
        user_id=22,
        question="What must the capstone report include?",
        file_name="capstone.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        file_bytes=buffer.getvalue(),
    )

    assert response.answer == "The project requires a report [1]."
    assert response.sources[0].source_type == "uploaded_file"
    assert response.sources[0].source_id is None
    assert response.sources[0].title == "capstone.docx"
    assert "UML class diagram" in response.sources[0].preview
    assert repository.replace_arguments is None
    assert repository.search_arguments is not None


@pytest.mark.anyio
async def test_chat_file_rejects_upload_above_byte_limit_before_extraction() -> None:
    service = RagChatService(
        session=object(),
        repository=FakeRagRepository([]),
        ai_client=FakeGeminiClient(),
    )

    with pytest.raises(ClassFlowError) as exc_info:
        await service.answer_class_question_with_file(
            classroom_id=10,
            user_id=22,
            question="Summarize this file",
            file_name="large.pdf",
            content_type="application/pdf",
            file_bytes=b"%PDF-" + b"x" * settings.RAG_CHAT_FILE_MAX_SIZE_BYTES,
        )

    assert exc_info.value.status_code == 413
    assert exc_info.value.detail["error_code"] == "CHAT_FILE_TOO_LARGE"


@pytest.mark.anyio
async def test_chat_file_rejects_unsupported_type() -> None:
    service = RagChatService(
        session=object(),
        repository=FakeRagRepository([]),
        ai_client=FakeGeminiClient(),
    )

    with pytest.raises(ClassFlowError) as exc_info:
        await service.answer_class_question_with_file(
            classroom_id=10,
            user_id=22,
            question="Read this file",
            file_name="notes.txt",
            content_type="text/plain",
            file_bytes=b"Private notes",
        )

    assert exc_info.value.status_code == 422
    assert exc_info.value.detail["error_code"] == "CHAT_FILE_TYPE_NOT_ALLOWED"


@pytest.mark.anyio
async def test_chat_file_rejects_user_without_approved_class_access() -> None:
    ai_client = FakeGeminiClient()
    service = RagChatService(
        session=object(),
        repository=FakeRagRepository([], has_class_access=False),
        ai_client=ai_client,
    )

    with pytest.raises(ClassFlowError) as exc_info:
        await service.answer_class_question_with_file(
            classroom_id=10,
            user_id=99,
            question="Summarize this file",
            file_name="notes.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            file_bytes=b"not-read-because-access-is-denied",
        )

    assert exc_info.value.status_code == 403
    assert exc_info.value.detail["error_code"] == "APPROVED_CLASS_MEMBERSHIP_REQUIRED"
    assert ai_client.embedded_documents is None
